import os
import sys
import io
import re
import argparse
from collections import Counter
import fitz
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')


DEFAULT_EXCLUDE_TEXTS = [
]

DEFAULT_REMOVE_LOGOS = True
DEFAULT_AUTO_FOOTER = True


PUA_BULLET_CHARS = {
    '\uf07d', '\uf0e3', '\uf06c', '\uf0d8', '\uf0a7', '\uf0b7', '\uf0fc',
    '\uf076', '\uf09f', '\uf077', '\uf0a8', '\uf0fe', '\uf02a', '\uf06e',
    '\uf075', '\uf0d9',
    '\u25b0', '\u25b1', '\u25aa', '\u25ab', '\u25ba', '\u25bb', '\u25b8', '\u25b9',
    '\u25a0', '\u25a1', '\u25cf', '\u25cb', '\u25c6', '\u25c7', '\u2022', '\u2023',
    '\u2043', '\u2219',
}

PUA_MATH_MAP = {
    '\uf02b': '+',
    '\uf02d': '−',
    '\uf0b1': '±',
    '\uf0b4': '×',
    '\uf0b8': '÷',
    '\uf0b9': '≠',
    '\uf0a3': '≤',
    '\uf0b3': '≥',
    '\uf0bb': '≈',
    '\uf0a5': '∞',
    '\uf0ae': '→',
    '\uf0ac': '←',
    '\uf0ab': '↔',
    '\uf0de': '⇒',
    '\uf0db': '⇔',
    '\uf0ce': '∈',
    '\uf0cf': '∉',
    '\uf0cc': '⊂',
    '\uf0cd': '⊆',
    '\uf0c7': '∩',
    '\uf0c8': '∪',
    '\uf0d9': '∧',
    '\uf0da': '∨',
    '\uf0d8': '¬',
    '\uf022': '∀',
    '\uf024': '∃',
    '\uf066': '∅',
    '\uf049': '⋂',
    '\uf055': '⋃',
    '\uf0e5': '∑',
    '\uf0d5': '∏',
    '\uf0be': '', 
}

SYSTEM_FONT_CANDIDATES = {
    'calibri': r'C:\Windows\Fonts\calibri.ttf',
    'calibri-regular': r'C:\Windows\Fonts\calibri.ttf',
    'calibrib': r'C:\Windows\Fonts\calibrib.ttf',
    'calibri-bold': r'C:\Windows\Fonts\calibrib.ttf',
    'calibri,bold': r'C:\Windows\Fonts\calibrib.ttf',
    'calibrii': r'C:\Windows\Fonts\calibrii.ttf',
    'calibri-italic': r'C:\Windows\Fonts\calibrii.ttf',
    'calibri,italic': r'C:\Windows\Fonts\calibrii.ttf',
    'arial': r'C:\Windows\Fonts\arial.ttf',
    'arialmt': r'C:\Windows\Fonts\arial.ttf',
    'arialbd': r'C:\Windows\Fonts\arialbd.ttf',
    'arial-boldmt': r'C:\Windows\Fonts\arialbd.ttf',
    'arial-bold': r'C:\Windows\Fonts\arialbd.ttf',
    'arial,bold': r'C:\Windows\Fonts\arialbd.ttf',
    'ariali': r'C:\Windows\Fonts\ariali.ttf',
    'arial-italicmt': r'C:\Windows\Fonts\ariali.ttf',
    'arial-italic': r'C:\Windows\Fonts\ariali.ttf',
    'arial,italic': r'C:\Windows\Fonts\ariali.ttf',
    'arialunicodems': r'C:\Windows\Fonts\arialuni.ttf',
    'timesnewroman': r'C:\Windows\Fonts\times.ttf',
    'timesnewromanpsmt': r'C:\Windows\Fonts\times.ttf',
    'timesnewromanps-boldmt': r'C:\Windows\Fonts\timesbd.ttf',
    'timesnewroman-bold': r'C:\Windows\Fonts\timesbd.ttf',
    'timesnewroman,bold': r'C:\Windows\Fonts\timesbd.ttf',
    'consolas': r'C:\Windows\Fonts\consola.ttf',
    'consolas-bold': r'C:\Windows\Fonts\consolab.ttf',
    'tahoma': r'C:\Windows\Fonts\tahoma.ttf',
    'tahoma-bold': r'C:\Windows\Fonts\tahomabd.ttf',
    'verdana': r'C:\Windows\Fonts\verdana.ttf',
    'verdana-bold': r'C:\Windows\Fonts\verdanab.ttf',
    'cambria': r'C:\Windows\Fonts\cambria.ttc',
}

SERBIAN_LATIN_CPS = {
    0x0106, 0x0107,  # Ć, ć
    0x010C, 0x010D,  # Č, č
    0x0110, 0x0111,  # Đ, đ
    0x0160, 0x0161,  # Š, š
    0x017D, 0x017E,  # Ž, ž
}


# POPRAVKA FONTOVA I SRPSKIH SLOVA (ToUnicode CMap Repair)
def _parse_tounicode_cmap(cmap_str):
    mapping = {}
    for block in re.findall(r'beginbfchar(.*?)endbfchar', cmap_str, re.S):
        for m in re.finditer(r'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>', block):
            cid = int(m.group(1), 16)
            uni_hex = m.group(2)
            if len(uni_hex) <= 4:
                mapping[cid] = int(uni_hex, 16)
    for block in re.findall(r'beginbfrange(.*?)endbfrange', cmap_str, re.S):
        for line in block.strip().splitlines():
            line = line.strip()
            if not line:
                continue
            m_arr = re.match(r'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*\[(.*)\]', line)
            if m_arr:
                start = int(m_arr.group(1), 16)
                end = int(m_arr.group(2), 16)
                vals = [int(x, 16) for x in re.findall(r'<([0-9A-Fa-f]+)>', m_arr.group(3))]
                for i, cid in enumerate(range(start, end + 1)):
                    if i < len(vals):
                        mapping[cid] = vals[i]
                continue
            m_rng = re.match(r'<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>', line)
            if m_rng:
                start = int(m_rng.group(1), 16)
                end = int(m_rng.group(2), 16)
                base = int(m_rng.group(3), 16)
                for i, cid in enumerate(range(start, end + 1)):
                    mapping[cid] = base + i
    return mapping


def _build_tounicode_stream(mapping):
    bfchars = [f"<{cid:04X}> <{uni:04X}>" for cid, uni in sorted(mapping.items()) if 0 <= uni <= 0xFFFF]
    chunks = []
    for i in range(0, len(bfchars), 100):
        chunk = bfchars[i:i + 100]
        chunks.append(f"{len(chunk)} beginbfchar\n" + "\n".join(chunk) + "\nendbfchar")
    return (
        "/CIDInit /ProcSet findresource begin\n"
        "12 dict begin\n"
        "begincmap\n"
        "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
        "/CMapName /Adobe-Identity-UCS def\n"
        "/CMapType 2 def\n"
        "1 begincodespacerange\n"
        "<0000> <FFFF>\n"
        "endcodespacerange\n" +
        "\n".join(chunks) + "\n"
        "endcmap\n"
        "CMapName currentdict /CMap defineresource pop\n"
        "end\n"
        "end\n"
    ).encode('latin1')


def fix_pdf_tounicode(doc):
    scan_codepoints = (
        list(range(0x0020, 0x0250)) +  # Osnovna latinica + srpska slova (č, ć, ž, š, đ, Č, Ć, Ž, Š, Đ)
        list(range(0x0370, 0x04FF)) +  # Grčki alfabet i ćirilica
        list(range(0x2000, 0x2300)) +  # Interpunkcija, strelice, matematički operatori
        list(range(0xF000, 0xF100))    # Private Use Area (Wingdings / Symbol)
    )

    for xref in range(1, doc.xref_length()):
        try:
            obj = doc.xref_object(xref)
        except Exception:
            continue
        if '/ToUnicode' not in obj:
            continue
        tu_m = re.search(r'/ToUnicode\s+(\d+)\s+0\s+R', obj)
        if not tu_m:
            continue
        tu_xref = int(tu_m.group(1))

        cid_xref = xref
        df_match = re.search(r'/DescendantFonts\s+(\d+)\s+0\s+R', obj)
        if df_match:
            try:
                df_obj = doc.xref_object(int(df_match.group(1)))
                m_cid = re.search(r'(\d+)\s+0\s+R', df_obj)
                if m_cid:
                    cid_xref = int(m_cid.group(1))
            except Exception:
                pass
        else:
            m2 = re.search(r'/DescendantFonts\s*\[\s*(\d+)\s+0\s+R', obj)
            if m2:
                cid_xref = int(m2.group(1))

        try:
            cid_obj = doc.xref_object(cid_xref)
        except Exception:
            continue

        ff2_xref = None
        fd_match = re.search(r'/FontDescriptor\s+(\d+)\s+0\s+R', cid_obj)
        if fd_match:
            try:
                fd_obj = doc.xref_object(int(fd_match.group(1)))
                ff2_m = re.search(r'/FontFile2\s+(\d+)\s+0\s+R', fd_obj)
                if ff2_m:
                    ff2_xref = int(ff2_m.group(1))
            except Exception:
                pass

        if not ff2_xref:
            continue

        try:
            ttf_bytes = doc.xref_stream(ff2_xref)
            f_emb = fitz.Font(fontbuffer=ttf_bytes)
        except Exception:
            continue

        gid_to_uni = {}
        for cp in scan_codepoints:
            try:
                g = f_emb.has_glyph(cp)
                if g > 0 and g not in gid_to_uni:
                    gid_to_uni[g] = cp
            except Exception:
                pass

        if not gid_to_uni and '/CIDToGIDMap /Identity' in cid_obj:
            base_m = re.search(r'/BaseFont\s+/([A-Za-z0-9+,\-_]+)', obj)
            if base_m:
                clean_font_name = re.sub(r'^[A-Z]{6}\+', '', base_m.group(1)).lower()
                sys_path = SYSTEM_FONT_CANDIDATES.get(clean_font_name)
                if sys_path and os.path.exists(sys_path):
                    try:
                        f_sys = fitz.Font(fontfile=sys_path)
                        for cp in scan_codepoints:
                            g = f_sys.has_glyph(cp)
                            if g > 0 and g not in gid_to_uni:
                                gid_to_uni[g] = cp
                    except Exception:
                        pass

        if not gid_to_uni:
            continue

        try:
            cmap_str = doc.xref_stream(tu_xref).decode('latin1', errors='ignore')
            tu_map = _parse_tounicode_cmap(cmap_str)
        except Exception:
            continue

        # Provera kompatibilnosti GID tabele sa postojećim ToUnicode mapiranjem:
        # 1) Ako u ToUnicode već postoje srpska slova i font se ne slaže sa njima, GID tabela nije ista (npr. druga verzija fonta).
        # 2) Najmanje 75% postojećih CID unosa mora da se poklapa sa GID tabelom.
        if tu_map:
            serbian_conflict = any(
                (u in SERBIAN_LATIN_CPS and gid_to_uni.get(cid) != u)
                for cid, u in tu_map.items()
            )
            if serbian_conflict:
                continue
            matches = sum(1 for cid, u in tu_map.items() if gid_to_uni.get(cid) == u)
            if matches / len(tu_map) < 0.75:
                continue

        needs_patch = False
        merged = dict(tu_map)
        for cid, true_uni in gid_to_uni.items():
            if cid in merged:
                old_uni = merged[cid]
                if old_uni != true_uni:
                    # Prepravljamo samo ako je u ToUnicode srpsko slovo pogrešno mapirano na ASCII znak (npr. ć -> d, Ć -> D)
                    # ili na nevalidan/zamenski karakter. Nikada ne gazimo već ispravna ne-ASCII slova!
                    if (0x0020 <= old_uni <= 0x007E and true_uni in SERBIAN_LATIN_CPS) or old_uni in (0, 0xFFFD):
                        merged[cid] = true_uni
                        needs_patch = True
            else:
                merged[cid] = true_uni

        if needs_patch:
            try:
                doc.update_stream(tu_xref, _build_tounicode_stream(merged))
            except Exception:
                pass


def clean_line_characters(line):
    if not line:
        return ""

    line = line.replace('\uf0be\uf0ae', '→').replace('\uf0be', '')

    stripped = line.lstrip()
    leading_ws = line[:len(line) - len(stripped)]

    if stripped and stripped[0] in PUA_BULLET_CHARS:
        rest = stripped[1:].lstrip()
        if not rest:
            return ""
        stripped = "• " + rest
        line = leading_ws + stripped

    out_chars = []
    for ch in line:
        if ch in PUA_MATH_MAP:
            out_chars.append(PUA_MATH_MAP[ch])
        elif ch in PUA_BULLET_CHARS:
            out_chars.append('• ')
        elif 0xF000 <= ord(ch) <= 0xF0FF:
            out_chars.append('• ')
        else:
            out_chars.append(ch)

    result = "".join(out_chars)
    result = re.sub(r'^(•\s*)+', '• ', result.strip()) if result.strip().startswith('•') else result
    return result


def _normalize_for_match(text):
    t = text.lower().strip()
    replacements = {
        'č': 'c', 'ć': 'c', 'ž': 'z', 'š': 's', 'đ': 'dj',
    }
    t_ascii = "".join(replacements.get(c, c) for c in t)
    t_clean = re.sub(r'[^\w\s]', ' ', t)
    t_clean = re.sub(r'\s+', ' ', t_clean).strip()
    t_ascii_clean = re.sub(r'[^\w\s]', ' ', t_ascii)
    t_ascii_clean = re.sub(r'\s+', ' ', t_ascii_clean).strip()
    return t_clean, t_ascii_clean


def is_bullet_start(text):
    t = text.strip()
    if not t:
        return False
    if t[0] in PUA_BULLET_CHARS or (0xF000 <= ord(t[0]) <= 0xF0FF):
        return True
    if re.match(r'^(\d+[\.\)]|[a-zA-Z][\.\)]|[-–—•])\s+', t):
        return True
    if re.match(r'^\d+\.$', t):
        return True
    return False


def get_page_text_lines(page):
    """
    Izvlači sve linije teksta sa stranice u vizuelnom (rotiranom) koordinatnom sistemu,
    uklanja duplirane drop-shadow kopije teksta i spaja horizontalne segmente koji se nalaze
    u istom vizuelnom redu (npr. bullet simbol + tekst ili reči razdvojene zbog poravnanja).
    """
    rot = page.rotation_matrix
    d = page.get_text("dict")
    raw_lines = []
    for b_idx, b in enumerate(d.get("blocks", [])):
        if b.get("type") != 0:
            continue
        for l in b.get("lines", []):
            r = fitz.Rect(l["bbox"]) * rot
            spans_text = []
            for s in l.get("spans", []):
                st = s.get("text", "")
                if st:
                    spans_text.append(st)
            txt = "".join(spans_text).strip()
            if not txt:
                continue
            raw_lines.append({
                "rect": r,
                "text": txt,
                "block_idx": b_idx,
            })

    # 1. Deduplikacija identičnih preklapajućih linija (npr. senka teksta u dijagramima)
    deduped = []
    for item in raw_lines:
        r1 = item["rect"]
        t1 = item["text"]
        is_dup = False
        for prev in deduped:
            if prev["text"] == t1:
                r2 = prev["rect"]
                inter = r1 & r2
                if not inter.is_empty:
                    ia = inter.width * inter.height
                    min_a = max(1.0, min(r1.width * r1.height, r2.width * r2.height))
                    if ia / min_a >= 0.45 or (abs(r1.x0 - r2.x0) <= 4 and abs(r1.y0 - r2.y0) <= 4):
                        is_dup = True
                        break
        if not is_dup:
            deduped.append(item)

    # 2. Spajanje horizontalnih segmenata u istom vizuelnom redu
    deduped.sort(key=lambda x: (round(x["rect"].y1 / 6.0), x["rect"].x0))
    merged = []
    for item in deduped:
        if not merged:
            merged.append(item)
            continue
        prev = merged[-1]
        r1, r2 = prev["rect"], item["rect"]
        y_ov = min(r1.y1, r2.y1) - max(r1.y0, r2.y0)
        min_h = max(1.0, min(r1.height, r2.height))
        x_gap = r2.x0 - r1.x1
        same_row = (y_ov / min_h >= 0.55) and (
            abs((r1.y0 + r1.y1) / 2 - (r2.y0 + r2.y1) / 2) <= 6.0
            or abs(r1.y1 - r2.y1) <= 5.0
        )
        can_join = False
        if same_row and x_gap >= -3.0:
            if prev["block_idx"] == item["block_idx"]:
                can_join = True
            elif (
                prev["text"] in PUA_BULLET_CHARS
                or re.match(r'^(\d+[\.\)]|[a-zA-Z][\.\)])$', prev["text"])
            ) and x_gap <= 55.0:
                can_join = True
        if can_join:
            sep = " " if (not prev["text"].endswith(" ") and not item["text"].startswith(" ")) else ""
            prev["text"] = prev["text"] + sep + item["text"]
            prev["rect"] = r1 | r2
        else:
            merged.append(item)

    return merged


def detect_repeated_slide_texts(doc):
    total_pages = len(doc)
    if total_pages < 3:
        return set()

    line_page_counts = Counter()
    for pno, page in enumerate(doc):
        ph = page.rect.height
        seen_on_page = set()
        for tl in get_page_text_lines(page):
            r = tl["rect"]
            in_margin = (r.y0 <= 0.18 * ph) or (r.y1 >= 0.82 * ph)
            if not in_margin:
                continue
            line = clean_line_characters(tl["text"]).strip()
            if not line or line.isdigit():
                continue
            norm, norm_ascii = _normalize_for_match(line)
            if len(norm) >= 3:
                seen_on_page.add(norm_ascii)
        for norm_ascii in seen_on_page:
            line_page_counts[norm_ascii] += 1

    threshold = max(3, int(total_pages * 0.30))
    repeated = {txt for txt, cnt in line_page_counts.items() if cnt >= threshold}
    return repeated


def should_exclude_line(line, y0, y1, page_height, pno, user_patterns, auto_repeated_set, pres_titles=None):
    t = line.strip()
    if not t:
        return True

    in_top_margin = (y0 <= 0.16 * page_height)
    in_bottom_margin = (y1 >= 0.82 * page_height) or (y0 >= 0.80 * page_height)
    in_margin_or_title_slide = in_top_margin or in_bottom_margin or (pno == 0)

    if (in_top_margin or in_bottom_margin) and re.match(r'^\d+(\s*/\s*\d+)?$', t):
        return True

    norm, norm_ascii = _normalize_for_match(t)
    if not norm_ascii:
        return True

    if norm_ascii in auto_repeated_set and in_margin_or_title_slide:
        return True

    if pno == 0 and auto_repeated_set:
        line_words = set(w for w in norm_ascii.split() if len(w) > 2)
        for rep in auto_repeated_set:
            rep_words = set(w for w in rep.split() if len(w) > 2)
            if len(rep_words) >= 2 and rep_words.issubset(line_words) and len(line_words) <= len(rep_words) + 3:
                return True

    if pres_titles and in_bottom_margin:
        for pt in pres_titles:
            _, pt_ascii = _normalize_for_match(pt)
            if pt_ascii and pt_ascii in norm_ascii and len(norm_ascii) <= len(pt_ascii) + 15:
                return True

    for pat in user_patterns:
        _, pat_ascii = _normalize_for_match(pat)
        if not pat_ascii:
            continue

        if norm_ascii == pat_ascii:
            return True

        if pat_ascii in norm_ascii:
            if in_margin_or_title_slide or len(norm_ascii) <= len(pat_ascii) + 25:
                return True

        if in_margin_or_title_slide:
            pat_words = [w for w in pat_ascii.split() if len(w) > 2]
            line_words = set(w for w in norm_ascii.split() if len(w) > 2)
            if len(pat_words) >= 2:
                matched = sum(1 for w in pat_words if w in line_words)
                if matched == len(pat_words) and len(line_words) <= len(pat_words) + 3:
                    return True

    return False


def detect_repeating_images(doc):
    if len(doc) < 2:
        return set(), set()
    xref_counts = Counter()
    sig_counts = Counter()
    for page in doc:
        seen_xrefs = set()
        seen_sigs = set()
        for im in page.get_images():
            xref = im[0]
            seen_xrefs.add(xref)
            try:
                base = doc.extract_image(xref)
                sig = (base['width'], base['height'], len(base['image']))
                seen_sigs.add(sig)
            except Exception:
                pass
        for x in seen_xrefs:
            xref_counts[x] += 1
        for s in seen_sigs:
            sig_counts[s] += 1

    threshold = 2
    rep_xrefs = {x for x, c in xref_counts.items() if c >= threshold}
    rep_sigs = {s for s, c in sig_counts.items() if c >= threshold}
    return rep_xrefs, rep_sigs


def detect_template_drawings(doc):
    if len(doc) < 3:
        return set()
    rect_counts = Counter()
    for page in doc:
        rot = page.rotation_matrix
        pw, ph = page.rect.width, page.rect.height
        seen = set()
        for d in page.get_drawings():
            r = d['rect'] * rot
            if r.width < 1.0:
                r.x1 = r.x0 + 1.0
            if r.height < 1.0:
                r.y1 = r.y0 + 1.0
            r = fitz.Rect(max(0, r.x0), max(0, r.y0), min(pw, r.x1), min(ph, r.y1))
            seen.add((round(r.x0 / 2.0), round(r.y0 / 2.0), round(r.x1 / 2.0), round(r.y1 / 2.0)))
        for k in seen:
            rect_counts[k] += 1
    return {k for k, c in rect_counts.items() if c >= 3}


def extract_slide_visual_regions(page, pno, doc, rep_xrefs, rep_sigs, tmpl_draws, text_lines, remove_logos=True):
    """
    Detektuje sve vizuelne regione (rasterske slike, složene slike iz segmenata i vektorske dijagrame)
    u vizuelnom (rotiranom) koordinatnom sistemu stranice.
    """
    rot = page.rotation_matrix
    pw, ph = page.rect.width, page.rect.height

    # 1. Validni rasterski regioni
    raster_entries = []
    for im in page.get_images():
        xref, smask = im[0], im[1]
        if remove_logos and xref in rep_xrefs:
            continue
        try:
            base = doc.extract_image(xref)
        except Exception:
            continue
        w, h = base["width"], base["height"]
        sig = (w, h, len(base["image"]))
        if remove_logos and sig in rep_sigs:
            continue

        # Filtriranje PowerPoint drop-shadow / glow bitmapa iza teksta (smask > 0 + mali broj boja + preklapanje sa tekstom)
        if smask > 0:
            try:
                pil_rgb = Image.open(io.BytesIO(base["image"])).convert("RGB")
                small = pil_rgb.resize((min(64, max(1, w)), min(64, max(1, h))), Image.Resampling.NEAREST)
                raw_b = small.tobytes()
                colors = set(raw_b[i:i + 3] for i in range(0, len(raw_b), 3))
                uniq_c = len(colors)
            except Exception:
                uniq_c = 999
            if uniq_c <= 35:
                is_shadow = False
                for r_raw in page.get_image_rects(xref):
                    r = r_raw * rot
                    for tl in text_lines:
                        inter = r & tl["rect"]
                        if not inter.is_empty:
                            ia = inter.width * inter.height
                            if ia / max(1.0, min(r.width * r.height, tl["rect"].width * tl["rect"].height)) >= 0.25:
                                is_shadow = True
                                break
                    if is_shadow:
                        break
                if is_shadow:
                    continue

        for r_raw in page.get_image_rects(xref):
            r = (r_raw * rot) & page.rect
            if r.is_empty:
                continue
            if (w <= 35 and h <= 35) or (r.width <= 25 and r.height <= 25):
                continue
            if (r.height <= 12 and r.width > 120) or (h <= 14 and w > 180):
                continue
            if (r.width <= 12 and r.height > 120) or (w <= 14 and h > 180):
                continue
            if r.y0 >= 0.88 * ph:
                continue
            if remove_logos:
                if r.width >= 0.85 * pw and r.height >= 0.85 * ph:
                    continue
                in_top_band = (r.y1 <= 0.28 * ph) and (r.y0 <= 0.16 * ph)
                in_left = (r.x1 <= 0.26 * pw)
                in_right = (r.x0 >= 0.74 * pw)
                is_emblem = (r.width <= 0.24 * pw) and (r.height <= 0.26 * ph)
                if in_top_band and (in_left or in_right) and is_emblem:
                    continue
                in_bot_band = (r.y0 >= 0.75 * ph) or (pno == 0 and r.y0 >= 0.55 * ph)
                if in_bot_band and (in_left or in_right) and is_emblem:
                    continue
                if pno == 0:
                    continue

            raster_entries.append({
                "rect": r,
                "w": w,
                "h": h,
                "xref": xref,
                "smask": smask,
            })

    # Deduplikacija preklapajućih rasterskih pravougaonika
    kept_rasters = []
    for item in raster_entries:
        r1 = item["rect"]
        overlapped = False
        for k_idx, k_item in enumerate(kept_rasters):
            r2 = k_item["rect"]
            inter = r1 & r2
            if not inter.is_empty:
                ia = inter.width * inter.height
                if ia / max(1.0, min(r1.width * r1.height, r2.width * r2.height)) > 0.55:
                    overlapped = True
                    if (r1.width * r1.height) > (r2.width * r2.height):
                        kept_rasters[k_idx] = item
                    break
        if not overlapped:
            kept_rasters.append(item)

    # 2. Detekcija klastera vektorskih dijagrama (preskače se naslovni slajd pno == 0)
    vector_clusters = []
    if pno > 0:
        cand_drawings = []
        for d in page.get_drawings():
            r = d["rect"] * rot
            if r.width < 1.0:
                r.x1 = r.x0 + 1.0
            if r.height < 1.0:
                r.y1 = r.y0 + 1.0
            if r.width >= 0.85 * pw and r.height >= 0.85 * ph:
                continue
            if d.get("fill") == (1.0, 1.0, 1.0) and d.get("color") is None:
                continue
            if (r.x0 < -5 or r.x1 > pw + 5) and r.width >= 0.85 * pw:
                continue
            r = fitz.Rect(max(0, r.x0), max(0, r.y0), min(pw, r.x1), min(ph, r.y1))
            if r.width <= 0 or r.height <= 0:
                continue
            if r.y1 <= 0.12 * ph or r.y0 >= 0.94 * ph:
                continue
            key = (round(r.x0 / 2.0), round(r.y0 / 2.0), round(r.x1 / 2.0), round(r.y1 / 2.0))
            if key in tmpl_draws:
                continue
            if r.width <= 14 and r.height <= 14:
                continue
            cand_drawings.append(r)

        # Izbacivanje izolovanih tankih horizontalnih linija (npr. podvlačenje reči u tekstu)
        filtered_drawings = []
        for i, r in enumerate(cand_drawings):
            if r.height <= 4.5 and r.width >= 20:
                has_neighbor = False
                exp = fitz.Rect(r.x0 - 18, r.y0 - 18, r.x1 + 18, r.y1 + 18)
                for j, r2 in enumerate(cand_drawings):
                    if i != j and exp.intersects(r2):
                        has_neighbor = True
                        break
                if not has_neighbor:
                    continue
            filtered_drawings.append(r)

        # Prepoznavanje običnih tekstualnih pasusa na slajdu (dugi redovi, bullet stavke i njihovi prelomljeni nastavci)
        prose_indices = set()
        for idx_tl, tl in enumerate(text_lines):
            if len(tl["text"]) > 26 or tl["rect"].width > 155 or is_bullet_start(tl["text"]):
                prose_indices.add(idx_tl)
        changed_p = True
        while changed_p:
            changed_p = False
            for idx_tl, tl in enumerate(text_lines):
                if idx_tl in prose_indices:
                    continue
                r_curr = tl["rect"]
                for idx_p in list(prose_indices):
                    r_prev = text_lines[idx_p]["rect"]
                    if tl["block_idx"] == text_lines[idx_p]["block_idx"] or (
                        abs(r_curr.x0 - r_prev.x0) <= 6.0 and -5.0 <= (r_curr.y0 - r_prev.y1) <= 10.0
                    ):
                        prose_indices.add(idx_tl)
                        changed_p = True
                        break

        bridge_labels = [
            tl["rect"] for idx_tl, tl in enumerate(text_lines)
            if idx_tl not in prose_indices
            and tl["rect"].y0 > 0.14 * ph
            and tl["rect"].y1 < 0.93 * ph
        ]

        n_d = len(filtered_drawings)
        visited = [False] * n_d
        for i in range(n_d):
            if visited[i]:
                continue
            visited[i] = True
            comp_indices = [i]
            comp_rect = fitz.Rect(filtered_drawings[i])
            changed = True
            while changed:
                changed = False
                exp_comp = fitz.Rect(comp_rect.x0 - 22, comp_rect.y0 - 22, comp_rect.x1 + 22, comp_rect.y1 + 22)
                for idx_c in comp_indices:
                    dr_c = filtered_drawings[idx_c]
                    near_d = fitz.Rect(dr_c.x0 - 12, dr_c.y0 - 12, dr_c.x1 + 12, dr_c.y1 + 12)
                    for bl in bridge_labels:
                        if near_d.intersects(bl):
                            exp_comp |= fitz.Rect(bl.x0 - 14, bl.y0 - 14, bl.x1 + 14, bl.y1 + 14)
                for j in range(n_d):
                    if not visited[j] and exp_comp.intersects(filtered_drawings[j]):
                        visited[j] = True
                        comp_indices.append(j)
                        comp_rect |= filtered_drawings[j]
                        changed = True

            touches_raster = any(
                fitz.Rect(comp_rect.x0 - 15, comp_rect.y0 - 15, comp_rect.x1 + 15, comp_rect.y1 + 15).intersects(re_item["rect"])
                for re_item in kept_rasters
            )
            min_drawings = 3 if touches_raster else 4
            if len(comp_indices) >= min_drawings and comp_rect.width >= 55 and comp_rect.height >= 24:
                expanded = fitz.Rect(comp_rect)
                comp_draw_rects = [filtered_drawings[idx_c] for idx_c in comp_indices]
                for _ in range(2):
                    for idx_tl, tl in enumerate(text_lines):
                        tr = tl["rect"]
                        if tr.y0 <= 0.13 * ph or tr.y1 >= 0.95 * ph:
                            continue
                        if is_bullet_start(tl["text"]):
                            continue
                        if re.match(r'^\d+(\s*/\s*\d+)?$', tl["text"].strip()) and (tr.y1 >= 0.85 * ph or tr.x0 >= 0.88 * pw):
                            continue
                        inter = tr & comp_rect
                        inside_ratio = (inter.width * inter.height) / max(1.0, tr.width * tr.height) if not inter.is_empty else 0.0
                        near_any_drawing = any(
                            fitz.Rect(dr_c.x0 - 14, dr_c.y0 - 16, dr_c.x1 + 14, dr_c.y1 + 16).intersects(tr)
                            for dr_c in comp_draw_rects
                        )
                        if inside_ratio >= 0.45 or (near_any_drawing and idx_tl not in prose_indices and len(tl["text"]) <= 25 and tr.width <= 145):
                            expanded |= tr
                            comp_draw_rects.append(tr)
                vector_clusters.append(expanded)

    # 3. Spajanje rasterskih regiona i vektorskih klastera u jedinstvene vizuelne regione
    regions = [{"rect": fitz.Rect(r["rect"]), "type": "raster", "rasters": [r]} for r in kept_rasters]
    for vc in vector_clusters:
        regions.append({"rect": fitz.Rect(vc), "type": "vector", "rasters": []})

    bullet_rects = [tl["rect"] for tl in text_lines if is_bullet_start(tl["text"])]

    merged_any = True
    while merged_any and len(regions) > 1:
        merged_any = False
        for i in range(len(regions)):
            for j in range(i + 1, len(regions)):
                r1, r2 = regions[i]["rect"], regions[j]["rect"]
                t1, t2 = regions[i]["type"], regions[j]["type"]
                inter = r1 & r2
                actually_overlaps = not inter.is_empty and (inter.width * inter.height) > 4.0
                is_aligned_strip = (
                    (abs(r1.x0 - r2.x0) <= 6 and abs(r1.x1 - r2.x1) <= 6 and min(abs(r1.y1 - r2.y0), abs(r2.y1 - r1.y0)) <= 6) or
                    (abs(r1.y0 - r2.y0) <= 6 and abs(r1.y1 - r2.y1) <= 6 and min(abs(r1.x1 - r2.x0), abs(r2.x1 - r1.x0)) <= 6)
                )
                exp1 = fitz.Rect(r1.x0 - 8, r1.y0 - 8, r1.x1 + 8, r1.y1 + 8)
                vec_raster_touch = (t1 != "raster" or t2 != "raster") and exp1.intersects(r2)

                if actually_overlaps or is_aligned_strip or vec_raster_touch:
                    union_r = r1 | r2
                    swallows_bullet = any(
                        union_r.intersects(br) and not r1.intersects(br) and not r2.intersects(br)
                        for br in bullet_rects
                    )
                    if not swallows_bullet:
                        regions[i]["rect"] = union_r
                        regions[i]["type"] = "composite"
                        regions[i]["rasters"].extend(regions[j]["rasters"])
                        regions.pop(j)
                        merged_any = True
                        break
            if merged_any:
                break

    # Proširenje rasterskih/složenih regiona ako ih presecaju pojedinačni vektorski okviri ili markeri (npr. highlight pravougaonici)
    if pno > 0:
        for reg in regions:
            rr = reg["rect"]
            exp_r = fitz.Rect(rr.x0 - 6, rr.y0 - 6, rr.x1 + 6, rr.y1 + 6)
            for d in page.get_drawings():
                dr = d["rect"] * rot
                if dr.width < 1.0:
                    dr.x1 = dr.x0 + 1.0
                if dr.height < 1.0:
                    dr.y1 = dr.y0 + 1.0
                if (dr.width >= 0.85 * pw and dr.height >= 0.85 * ph) or (d.get("fill") == (1.0, 1.0, 1.0) and d.get("color") is None):
                    continue
                dr = fitz.Rect(max(0, dr.x0), max(0, dr.y0), min(pw, dr.x1), min(ph, dr.y1))
                if dr.width <= 0 or dr.height <= 0:
                    continue
                if dr.y1 <= 0.12 * ph or dr.y0 >= 0.94 * ph:
                    continue
                key = (round(dr.x0 / 2.0), round(dr.y0 / 2.0), round(dr.x1 / 2.0), round(dr.y1 / 2.0))
                if key in tmpl_draws:
                    continue
                if dr.width <= 14 and dr.height <= 14:
                    continue
                if exp_r.intersects(dr) and dr.width <= 0.85 * pw and dr.height <= 0.85 * ph:
                    union_r = reg["rect"] | dr
                    if not any(union_r.intersects(br) and not reg["rect"].intersects(br) for br in bullet_rects):
                        reg["rect"] = union_r
                        reg["type"] = "composite"

    regions.sort(key=lambda reg: (round(reg["rect"].y0 / 20.0), reg["rect"].x0))
    return regions


def insert_image(doc, img_bytes, w, h):
    try:
        pil_img = Image.open(io.BytesIO(img_bytes))
        pw, ph = pil_img.size

        # Konverzija u standardni PNG u memoriji kako python-docx nikada ne bi bacio UnrecognizedImageError
        if pil_img.mode not in ("RGB", "RGBA", "L"):
            pil_img = pil_img.convert("RGBA" if "A" in pil_img.mode else "RGB")
        png_stream = io.BytesIO()
        pil_img.save(png_stream, format="PNG")
        png_stream.seek(0)

        max_w = 5.6
        max_h = 4.5
        aspect = pw / max(1, ph)

        if pw < 250 and ph < 250:
            disp_w = min(2.5, pw / 96.0)
        elif aspect > 1.2:
            disp_w = min(max_w, pw / 120.0)
            if disp_w < 3.5:
                disp_w = 3.5
        else:
            disp_w = min(max_w, max_h * aspect)
            if disp_w < 3.0:
                disp_w = 3.0

        disp_w = min(disp_w, max_w)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True

        run = p.add_run()
        run.add_picture(png_stream, width=Inches(disp_w))
    except Exception as e:
        print(f"      [Upozorenje] Greška pri umetanju slike: {e}")


def convert_pdf_to_docx(
    pdf_path,
    out_docx_path=None,
    exclude_texts=None,
    remove_logos=True,
    auto_footer=True,
    master_doc=None
):
    if not out_docx_path:
        out_docx_path = os.path.splitext(pdf_path)[0] + ".docx"

    user_patterns = list(DEFAULT_EXCLUDE_TEXTS)
    if exclude_texts:
        for item in exclude_texts:
            if item and item.strip():
                user_patterns.append(item.strip())

    doc_docx = docx.Document() if master_doc is None else master_doc
    if master_doc is None:
        for s in doc_docx.sections:
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.left_margin = Inches(1)
            s.right_margin = Inches(1)

    base_name = os.path.splitext(os.path.basename(pdf_path))[0]
    pres_titles = [
        base_name.strip(),
        re.sub(r'^(\d+\.|\w+\s*-\s*\d+\s*-\s*|[A-Za-z]+\s+\d+)\s*', '', base_name).strip()
    ]

    title_p = doc_docx.add_paragraph()
    title_p.paragraph_format.space_before = Pt(14)
    title_p.paragraph_format.space_after = Pt(8)
    t_run = title_p.add_run(base_name)
    t_run.font.name = "Arial"
    t_run.font.size = Pt(16)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 44, 89)

    pdf_doc = fitz.open(pdf_path)

    fix_pdf_tounicode(pdf_doc)

    auto_repeated_set = detect_repeated_slide_texts(pdf_doc) if auto_footer else set()
    rep_xrefs, rep_sigs = detect_repeating_images(pdf_doc) if remove_logos else (set(), set())
    template_drawing_rects = detect_template_drawings(pdf_doc)

    total_pages = len(pdf_doc)
    images_count = 0

    for pno in range(total_pages):
        page = pdf_doc[pno]
        pw, ph = page.rect.width, page.rect.height

        slide_h = doc_docx.add_paragraph()
        slide_h.paragraph_format.space_before = Pt(16)
        slide_h.paragraph_format.space_after = Pt(4)
        slide_h.paragraph_format.keep_with_next = True
        h_run = slide_h.add_run(f"--- Slajd {pno + 1} ---")
        h_run.font.name = "Arial"
        h_run.font.size = Pt(12)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(31, 78, 121)

        text_lines = get_page_text_lines(page)
        visual_regions = extract_slide_visual_regions(
            page, pno, pdf_doc,
            rep_xrefs, rep_sigs, template_drawing_rects,
            text_lines, remove_logos=remove_logos
        )

        # Sortiranje linija teksta: leva kolona/glavni tok pa desna kolona, odozgo nadole
        text_lines_sorted = sorted(
            text_lines,
            key=lambda tl: (
                1 if tl["rect"].x0 >= 0.55 * pw else 0,
                round(tl["rect"].y0 / 4.0),
                round(tl["rect"].x0, 1)
            )
        )

        for tl in text_lines_sorted:
            tr = tl["rect"]
            y0, y1 = tr.y0, tr.y1
            raw_line = tl["text"]

            # Preskoči tekst koji se već nalazi unutar izvučene slike ili vektorskog dijagrama
            # (osim ako je u pitanju glavni naslov slajda na vrhu ili bullet stavka)
            is_slide_title = (y0 <= 0.14 * ph) and (tr.x0 <= 0.50 * pw)
            if visual_regions and not is_slide_title and not is_bullet_start(raw_line):
                in_visual_region = False
                for reg in visual_regions:
                    rr = fitz.Rect(
                        reg["rect"].x0 - 4,
                        reg["rect"].y0 - 4,
                        reg["rect"].x1 + 4,
                        reg["rect"].y1 + 4
                    )
                    inter = tr & rr
                    if not inter.is_empty:
                        ia = inter.width * inter.height
                        if ia / max(1.0, tr.width * tr.height) >= 0.45:
                            in_visual_region = True
                            break
                if in_visual_region:
                    continue

            cleaned_line = clean_line_characters(raw_line).rstrip()
            if not cleaned_line.strip():
                continue

            if should_exclude_line(
                cleaned_line, y0, y1, ph, pno,
                user_patterns, auto_repeated_set, pres_titles
            ):
                continue

            p = doc_docx.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(cleaned_line)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(34, 34, 34)

        # Renderovanje i umetanje svih vizuelnih regiona direktno sa slajda (čuva prozirnost / SMask, rotaciju i spaja segmente)
        for reg in visual_regions:
            rr = reg["rect"]
            clip_r = fitz.Rect(
                max(0, rr.x0 - 4),
                max(0, rr.y0 - 4),
                min(pw, rr.x1 + 4),
                min(ph, rr.y1 + 4)
            )
            if clip_r.is_empty or clip_r.width < 5 or clip_r.height < 5:
                continue
            pix = page.get_pixmap(clip=clip_r, dpi=200)
            insert_image(doc_docx, pix.tobytes("png"), pix.width, pix.height)
            images_count += 1

    if master_doc is None:
        doc_docx.save(out_docx_path)
        print(f"Konvertovano: {os.path.basename(pdf_path)} -> {os.path.basename(out_docx_path)} ({total_pages} slajdova, {images_count} slika)")
    return out_docx_path, total_pages, images_count


def main():
    parser = argparse.ArgumentParser(
        description="Konverter .pdf u .docx sa filtriranjem futera i logoa."
    )
    parser.add_argument("files", nargs="*", help="PDF fajlovi ili folderi za obradu (podrazumevano: trenutni folder).")
    parser.add_argument(
        "-e", "--exclude",
        nargs="+",
        default=[],
        help="Tekstovi koje treba izbaciti sa slajdova"
    )
    parser.add_argument(
        "--keep-logos",
        action="store_true",
        help="Zadrži logoe (podrazumevano se izbacuju)"
    )
    parser.add_argument(
        "--no-auto-footer",
        action="store_true",
        help="Isključi automatsku detekciju ponavljajućeg teksta u hederu/futeru."
    )
    parser.add_argument(
        "--interactive", "-i",
        action="store_true",
        help="Interaktivno pitaj za tekstove koje treba izbaciti pre pokretanja."
    )
    parser.add_argument("--all", action="store_true", help="Obradi sve PDF fajlove u folderu.")
    parser.add_argument("--master", action="store_true", help="Kreiraj i objedinjeni Word fajl.")
    args = parser.parse_args()

    exclude_list = list(DEFAULT_EXCLUDE_TEXTS)
    for item in args.exclude:
        for sub in re.split(r'[;,]', item):
            if sub.strip():
                exclude_list.append(sub.strip())

    remove_logos = DEFAULT_REMOVE_LOGOS and (not args.keep_logos)
    auto_footer = DEFAULT_AUTO_FOOTER and (not args.no_auto_footer)

    if args.interactive:
        print("=== PODEŠAVANJA KONVERZIJE PDF -> DOCX ===")
        raw_ex = input("Unesite tekstove za izbegavanje (odvojene zarezom, Enter za preskakanje): ").strip()
        if raw_ex:
            for sub in re.split(r'[;,]', raw_ex):
                if sub.strip():
                    exclude_list.append(sub.strip())
        raw_logo = input("Izbaciti logoe? [y/n]: ").strip().lower()
        if raw_logo in ('n', 'ne', 'no'):
            remove_logos = False

    target_files = []
    search_dir = "."
    if not args.files or args.all:
        pdfs = [f for f in os.listdir('.') if f.lower().endswith('.pdf')]
        target_files = pdfs
    else:
        for f in args.files:
            if os.path.isdir(f):
                search_dir = f
                for entry in os.listdir(f):
                    if entry.lower().endswith('.pdf'):
                        target_files.append(os.path.join(f, entry))
            elif os.path.exists(f) and f.lower().endswith('.pdf'):
                target_files.append(f)

    def sort_key(s):
        base = os.path.basename(s)
        m = re.search(r'(\d+)', base)
        return (int(m.group(1)) if m else 999, base.lower())

    target_files = sorted(target_files, key=sort_key)

    if not target_files:
        print("Nisu pronađeni PDF fajlovi za obradu.")
        return

    print(f"Pronađeno {len(target_files)} PDF fajlova za obradu.")
    if exclude_list:
        print(f"Izbacivanje zadatih tekstova: {exclude_list}")
    print(f"Automatsko čišćenje ponavljajućih futera/hedera: {'DA' if auto_footer else 'NE'}")
    print(f"Izbacivanje logoa i ukrasnih pozadina: {'DA' if remove_logos else 'NE'}\n")

    for f in target_files:
        convert_pdf_to_docx(
            f,
            exclude_texts=exclude_list,
            remove_logos=remove_logos,
            auto_footer=auto_footer
        )

    if args.master and len(target_files) > 1:
        print("\nKreiranje objedinjenog fajla...")
        master_doc = docx.Document()
        for s in master_doc.sections:
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.left_margin = Inches(1)
            s.right_margin = Inches(1)
        for idx, f in enumerate(target_files):
            if idx > 0:
                master_doc.add_page_break()
            convert_pdf_to_docx(
                f,
                exclude_texts=exclude_list,
                remove_logos=remove_logos,
                auto_footer=auto_footer,
                master_doc=master_doc
            )
        m_name = os.path.join(search_dir, "Objedinjeno.docx")
        master_doc.save(m_name)
        print(f"Kreiran objedinjeni fajl: {m_name}")


if __name__ == '__main__':
    main()